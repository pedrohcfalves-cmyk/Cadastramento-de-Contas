"""Estatísticas da pesquisa: coleta, gravação, junção de computadores e cálculo."""

import json
import time

import pytest

from sigef import estatisticas as est
from sigef.fluxos import Relator, executar_lotes, executar_planilha
from sigef.modelos import Lote, Resposta
from sigef.planilha import Planilha

from excel_falso import Aba, Pasta


def _sessao(itens, computador="PC1", usuario="ana", inicio="2026-10-09T08:00:00"):
    return {"tipo": "sessao", "id": f"{computador}-{inicio}-{len(itens)}",
            "computador": computador, "usuario": usuario, "inicio": inicio,
            "itens": [dict(zip(("status", "caminho", "segundos", "espera"), i)) for i in itens]}


def test_calculo_segue_a_metodologia():
    sessoes = [
        _sessao([("OK", "novo", 20, 5), ("OK", "existente", 10, 0), ("JA", "existente", 8, 0)]),
        _sessao([("ERRO", "novo", 12, 0), ("ERRO", "validacao", 0, 0), ("OK", "novo", 18, 0)],
                computador="PC2", usuario="joao", inicio="2026-10-10T09:00:00"),
    ]
    manuais = [{"tipo": "manual", "caminho": "novo", "segundos": 240},
               {"tipo": "manual", "caminho": "novo", "segundos": 360}]
    ind = est.calcular(sessoes, manuais, {"existente": 120})

    assert ind.feitos == 3
    assert ind.ok_por_caminho == {"novo": 2, "existente": 1}
    assert ind.media_manual == {"novo": 300, "existente": 120}      # medido / estimativa
    assert ind.fonte_manual("novo") == "medido (2 cronometragens)"
    assert ind.fonte_manual("existente") == "estimativa"
    assert ind.segundos_manual_equivalente == 2 * 300 + 120
    assert ind.segundos_automacao == 20 + 10 + 8 + 12 + 18          # validação fica fora
    assert ind.segundos_economizados == 720 - 68
    assert ind.media_automacao == pytest.approx(48 / 3)
    assert ind.segundos_espera == 5
    assert ind.taxa_sucesso == pytest.approx(4 / 6)
    assert ind.computadores == 2 and ind.usuarios == 2
    assert [d["dia"] for d in ind.por_dia] == ["2026-10-09", "2026-10-10"]
    assert ind.por_computador[1]["computador"] == "PC2"
    assert ind.vezes_mais_rapido == pytest.approx(240 / 16)


def test_sem_dados_nao_quebra():
    ind = est.calcular([], [])
    assert ind.feitos == 0 and ind.vezes_mais_rapido == 0 and ind.taxa_sucesso == 0
    assert ind.media_manual_geral == pytest.approx((300 + 180) / 2)


def test_grava_local_e_na_pasta_da_pesquisa(tmp_path, estatisticas_isoladas):
    compartilhada = tmp_path / "rede" / "pesquisa"
    est.gravar_ajustes({"pasta_pesquisa": str(compartilhada), "estimativa": {}})

    sessao = est.Sessao("lotes", "auto")
    sessao.registrar("OK", "novo", 15.0, 3.0)
    sessao.registrar("ERRO", "novo", 9.0, 0.0, "CPF 52998224725 inválido na Receita")
    est.salvar_sessao(sessao)
    est.salvar_tempo_manual("existente", 150)

    local = est.arquivo_local()
    linhas = [json.loads(l) for l in local.read_text(encoding="utf-8").splitlines()]
    assert [l["tipo"] for l in linhas] == ["sessao", "manual"]
    assert linhas[0]["itens"][0] == {"status": "OK", "caminho": "novo", "segundos": 12.0,
                                     "espera": 3.0, "motivo": ""}
    assert linhas[0]["itens"][1]["motivo"] == "CPF # inválido na Receita"   # sem o CPF
    assert linhas[0]["fim"] and linhas[0]["computador"] and linhas[0]["versao_app"]
    assert (compartilhada / local.name).read_text(encoding="utf-8") == local.read_text(encoding="utf-8")


def test_junta_computadores_sem_repetir(tmp_path, estatisticas_isoladas):
    compartilhada = tmp_path / "rede"
    compartilhada.mkdir()
    est.gravar_ajustes({"pasta_pesquisa": str(compartilhada), "estimativa": {}})

    sessao = est.Sessao("planilha", "auto")
    sessao.registrar("OK", "novo", 10)
    est.salvar_sessao(sessao)  # vai para o local E para a rede: não pode contar 2 vezes

    outro = tmp_path / "pen-drive" / "PC9__maria.jsonl"
    outro.parent.mkdir()
    outro.write_text(json.dumps(_sessao([("OK", "existente", 7, 0)], "PC9", "maria")) + "\n",
                     encoding="utf-8")
    assert est.importar_arquivos([outro]) == 1

    assert len(est.carregar(est.LOCAL)[0]) == 1
    sessoes, _ = est.carregar(est.TODOS)
    assert len(sessoes) == 2
    assert est.calcular(sessoes, []).computadores == 2


def test_ini_ao_lado_do_programa_define_a_pasta(tmp_path):
    (tmp_path / "CadastroSIGEF.ini").write_text("[pesquisa]\npasta = \\\\srv\\pesquisa\n",
                                                encoding="utf-8")
    assert est.ler_ajustes()["pasta_pesquisa"] == "\\\\srv\\pesquisa"


def test_formatar_duracao():
    assert est.formatar_duracao(42) == "42 s"
    assert est.formatar_duracao(200) == "3 min 20 s"
    assert est.formatar_duracao(3725) == "1 h 2 min"
    assert est.formatar_duracao(-90) == "−1 min 30 s"


# ---------------------------------------------------------- com o SIGEF simulado

def test_execucao_registra_tempos_e_caminhos(navegador):
    def perguntar(lote, acao, editavel):
        time.sleep(0.3)  # pessoa conferindo
        return Resposta.SIM_TODOS

    sessao = est.Sessao("lotes", "manual")
    executar_lotes(navegador, [
        Lote("71428793860", "NOVO TESTE", "001", "10", "20"),
        Lote("71428793860", "NOVO TESTE", "104", "10", "21"),
    ], Relator(), perguntar, sessao)

    novo, existente = sessao.itens
    assert (novo.status, novo.caminho) == ("OK", "novo")
    assert (existente.status, existente.caminho) == ("OK", "existente")
    assert novo.espera >= 0.3 and novo.segundos > 0
    assert sessao.fim and not sessao.interrompida


def test_planilha_marca_validacao(navegador):
    aba = Aba("A", [["SERVIDOR", "CPF", "BANCO", "AGÊNCIA", "CONTA"],
                    ["SEM CONTA", "71428793860", "001", "1", ""]])
    planilha = Planilha(Pasta("p.xlsx", [aba]), aba)
    colunas, primeira, _ = planilha.colunas_sugeridas()
    sessao = est.Sessao("planilha", "auto")
    executar_planilha(navegador, planilha, colunas, primeira, 2, "J", Relator(), None, sessao)
    assert [(i.status, i.caminho) for i in sessao.itens] == [("ERRO", "validacao")]


def test_relatorio_gera_os_tres_arquivos(tmp_path):
    from openpyxl import load_workbook

    from sigef import relatorio

    for itens in ([("OK", "novo", 20, 2), ("JA", "existente", 9, 0)],
                  [("OK", "existente", 11, 0), ("ERRO", "novo", 13, 0)]):
        sessao = est.Sessao("lotes", "auto")
        for status, caminho, segundos, espera in itens:
            sessao.registrar(status, caminho, segundos + espera, espera, "timeout")
        est.salvar_sessao(sessao)
    est.salvar_tempo_manual("novo", 280)

    pasta = relatorio.gerar(est.LOCAL, tmp_path / "saida")
    assert {p.name for p in pasta.iterdir()} == {"dados_pesquisa.xlsx", "relatorio.html",
                                                 "dados_brutos.jsonl"}

    livro = load_workbook(pasta / "dados_pesquisa.xlsx")
    assert livro.sheetnames == ["Resumo", "Sessões", "Itens", "Tempos manuais",
                                "Por computador", "Por dia", "Motivos de erro", "Metodologia"]
    assert livro["Sessões"].max_row == 3 and livro["Itens"].max_row == 5
    valores = {livro["Resumo"].cell(r, 1).value: livro["Resumo"].cell(r, 2).value
               for r in range(5, 30)}
    assert valores["Cadastros concluídos"] == 2
    assert valores["Tempo economizado"] == est.formatar_duracao(280 + 180 - (20 + 9 + 11 + 13))
    assert len(livro["Resumo"]._charts) == 2

    texto = (pasta / "relatorio.html").read_text(encoding="utf-8")
    assert "cadastros concluídos" in texto and "Como os números foram calculados" in texto
    assert len((pasta / "dados_brutos.jsonl").read_text(encoding="utf-8").splitlines()) == 3


def test_formatar_duracao_curta():
    assert est.formatar_duracao(1190, curto=True) == "20 min"
    assert est.formatar_duracao(200, curto=True) == "3 min 20 s"
    assert est.formatar_vezes(19.87) == "19,9×" and est.formatar_data("2026-10-09") == "09/10/2026"
