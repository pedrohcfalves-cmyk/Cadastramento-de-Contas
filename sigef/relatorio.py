"""Relatório da pesquisa: planilha Excel completa, página para apresentar e dados brutos.

Gera uma pasta com:
    dados_pesquisa.xlsx  resumo com gráfico, sessões, itens, tempos manuais,
                         por computador, por dia, motivos de erro e metodologia
    relatorio.html       resumo visual, pronto para projetar ou imprimir em PDF
    dados_brutos.jsonl   todos os registros, para juntar com outros computadores
"""

import html
import json
import sys
from datetime import datetime
from pathlib import Path

from . import __version__
from . import estatisticas as est
from .estatisticas import (
    EXISTENTE,
    NOVO,
    ROTULO_CAMINHO,
    formatar_data,
    formatar_duracao,
    formatar_vezes,
)

STATUS = {"OK": "Cadastrado", "JA": "Já tinha cadastro", "ERRO": "Erro", "PULADO": "Pulado"}
CAMINHOS = {NOVO: "Credor novo + conta", EXISTENTE: "Só a conta", est.VALIDACAO: "Recusado antes do SIGEF"}

METODOLOGIA = [
    "Cada execução da automação (sessão) é registrada automaticamente, com o tempo de cada "
    "pessoa processada. Não são gravados CPF, nome nem dados bancários.",
    "Tempo da automação: soma do tempo de máquina de TODOS os itens enviados ao SIGEF, "
    "inclusive os que deram erro ou já estavam cadastrados. O tempo de uma pessoa "
    "conferindo os dados na tela é medido à parte (coluna 'espera') e não entra nessa soma.",
    "Tempo manual: média das cronometragens feitas no próprio programa (botão 'Cronometrar "
    "um cadastro manual'), separada por tipo: credor novo + conta, ou só a conta. Sem "
    "cronometragens, usa a estimativa informada na tela de Estatísticas.",
    "Tempo manual equivalente: cadastros concluídos (situação 'Cadastrado') × tempo manual "
    "médio do mesmo tipo. Itens com erro ou já cadastrados não contam como economia "
    "(cálculo conservador).",
    "Tempo economizado = tempo manual equivalente − tempo da automação.",
    "Vezes mais rápido = tempo manual médio ÷ tempo médio da automação por cadastro concluído.",
]


def _documentos() -> Path:
    """Pasta Documentos de verdade (no Windows ela pode estar no OneDrive)."""
    if sys.platform == "win32":
        import ctypes

        caminho = ctypes.create_unicode_buffer(260)
        if ctypes.windll.shell32.SHGetFolderPathW(None, 5, None, 0, caminho) == 0:  # CSIDL_PERSONAL
            return Path(caminho.value)

    documentos = Path.home() / "Documents"
    return documentos if documentos.is_dir() else Path.home()


def pasta_padrao() -> Path:
    return _documentos() / "Relatorios Cadastro SIGEF"


def _minutos(segundos: float) -> float:
    return round(segundos / 60, 2)


def _linhas_sessoes(sessoes: list, media_manual: dict) -> list:
    linhas = []

    for s in sessoes:
        itens = s.get("itens", [])
        conta = {k: sum(i.get("status") == k for i in itens) for k in STATUS}
        manual, automacao = est.economia_de(itens, media_manual)
        linhas.append([
            s.get("inicio", "")[:10], s.get("inicio", "")[11:19], s.get("fim", "")[11:19],
            s.get("computador", ""), s.get("usuario", ""), s.get("versao_app", ""),
            s.get("interface", ""), s.get("modo", ""), s.get("conferencia", ""),
            len(itens), conta["OK"], conta["JA"], conta["ERRO"], conta["PULADO"],
            sum(i.get("status") == "OK" and i.get("caminho") == NOVO for i in itens),
            sum(i.get("status") == "OK" and i.get("caminho") == EXISTENTE for i in itens),
            round(automacao, 1), round(sum(i.get("espera", 0) for i in itens), 1),
            round(manual, 1), round(manual - automacao, 1),
            round(s.get("segundos_sessao", 0), 1),
            "sim" if s.get("interrompida") else "não", s.get("id", ""),
        ])

    return linhas


def _indicadores_tabela(ind: est.Indicadores) -> list:
    return [
        ("Período", f"{formatar_data(ind.primeira_data)} a {formatar_data(ind.ultima_data)}"
                    if ind.sessoes else "-"),
        ("Sessões (execuções)", ind.sessoes),
        ("Computadores", ind.computadores),
        ("Usuários", ind.usuarios),
        ("Pessoas processadas", ind.processados),
        ("Cadastros concluídos", ind.feitos),
        ("   credor novo + conta", ind.ok_por_caminho[NOVO]),
        ("   só a conta", ind.ok_por_caminho[EXISTENTE]),
        ("Já tinham cadastro", ind.contagem["JA"]),
        ("Com erro", ind.contagem["ERRO"]),
        ("Pulados na conferência", ind.contagem["PULADO"]),
        ("Taxa de sucesso", f"{ind.taxa_sucesso:.0%}"),
        ("Tempo manual médio por cadastro", formatar_duracao(ind.media_manual_geral)),
        (f"   {ROTULO_CAMINHO[NOVO]}",
         f"{formatar_duracao(ind.media_manual[NOVO])} ({ind.fonte_manual(NOVO)})"),
        (f"   {ROTULO_CAMINHO[EXISTENTE]}",
         f"{formatar_duracao(ind.media_manual[EXISTENTE])} ({ind.fonte_manual(EXISTENTE)})"),
        ("Tempo médio da automação por cadastro", formatar_duracao(ind.media_automacao)),
        ("Vezes mais rápido", formatar_vezes(ind.vezes_mais_rapido)),
        ("Tempo manual equivalente", formatar_duracao(ind.segundos_manual_equivalente)),
        ("Tempo total da automação", formatar_duracao(ind.segundos_automacao)),
        ("Tempo economizado", formatar_duracao(ind.segundos_economizados)),
        ("Tempo de pessoas conferindo (à parte)", formatar_duracao(ind.segundos_espera)),
    ]


# ================================================================ Excel

def _gerar_excel(destino: Path, ind: est.Indicadores, sessoes: list, manuais: list) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    titulo_fonte = Font(bold=True, size=14, color="0B3D91")
    cabecalho_fonte = Font(bold=True, color="FFFFFF")
    cabecalho_fundo = PatternFill("solid", fgColor="0B3D91")

    def tabela(aba, cabecalho: list, linhas: list, linha_inicial: int = 1) -> None:
        for coluna, texto in enumerate(cabecalho, 1):
            celula = aba.cell(row=linha_inicial, column=coluna, value=texto)
            celula.font, celula.fill = cabecalho_fonte, cabecalho_fundo
            celula.alignment = Alignment(wrap_text=True, vertical="center")

        for deslocamento, linha in enumerate(linhas, 1):
            for coluna, valor in enumerate(linha, 1):
                aba.cell(row=linha_inicial + deslocamento, column=coluna, value=valor)

        for coluna, texto in enumerate(cabecalho, 1):
            largura = max([len(str(texto))] + [len(str(l[coluna - 1])) for l in linhas[:200]])
            aba.column_dimensions[get_column_letter(coluna)].width = min(max(largura + 2, 10), 48)

        aba.freeze_panes = aba.cell(row=linha_inicial + 1, column=1)

    livro = Workbook()

    # ---- Resumo
    resumo = livro.active
    resumo.title = "Resumo"
    resumo["A1"] = "Pesquisa de implantação — Cadastro de Credores no SIGEF"
    resumo["A1"].font = titulo_fonte
    resumo["A2"] = f"Gerado em {datetime.now():%d/%m/%Y %H:%M} · versão {__version__}"
    tabela(resumo, ["Indicador", "Valor"], _indicadores_tabela(ind), linha_inicial=4)
    for (celula,) in resumo.iter_rows(min_row=5, min_col=2, max_col=2):
        celula.alignment = Alignment(horizontal="left")
    resumo.column_dimensions["A"].width = 44
    resumo.column_dimensions["B"].width = 40

    # Dados do gráfico: minutos por cadastro e totais
    base = 4 + len(_indicadores_tabela(ind)) + 3
    resumo.cell(row=base, column=1, value="Comparação (minutos)").font = Font(bold=True)
    dados_grafico = [
        ("", "Manual", "Com a automação"),
        ("Por cadastro", _minutos(ind.media_manual_geral), _minutos(ind.media_automacao)),
        ("Total do período", _minutos(ind.segundos_manual_equivalente),
         _minutos(ind.segundos_automacao)),
    ]
    for deslocamento, linha in enumerate(dados_grafico, 1):
        for coluna, valor in enumerate(linha, 1):
            resumo.cell(row=base + deslocamento, column=coluna, value=valor)

    for linha_grafico, titulo in ((base + 2, "Tempo por cadastro (min)"),
                                  (base + 3, "Tempo total do período (min)")):
        grafico = BarChart()
        grafico.type = "bar"
        grafico.title = titulo
        grafico.legend = None
        grafico.y_axis.title = "minutos"
        valores = Reference(resumo, min_col=2, max_col=3, min_row=linha_grafico,
                            max_row=linha_grafico)
        grafico.add_data(valores, from_rows=True, titles_from_data=False)
        grafico.set_categories(Reference(resumo, min_col=2, max_col=3, min_row=base + 1))
        grafico.height, grafico.width = 6, 16
        resumo.add_chart(grafico, f"D{4 if linha_grafico == base + 2 else 18}")

    # ---- Sessões
    tabela(livro.create_sheet("Sessões"), [
        "Data", "Início", "Fim", "Computador", "Usuário", "Versão", "Interface", "Modo",
        "Conferência", "Itens", "Cadastrados", "Já tinham", "Erros", "Pulados",
        "Concluídos: credor novo", "Concluídos: só conta", "Automação (s)", "Espera pessoas (s)",
        "Manual equivalente (s)", "Economia (s)", "Duração da sessão (s)", "Interrompida", "ID",
    ], _linhas_sessoes(sessoes, ind.media_manual))

    # ---- Itens
    itens = [
        [s.get("inicio", "")[:10], s.get("computador", ""), s.get("usuario", ""), numero,
         STATUS.get(i.get("status"), i.get("status")), CAMINHOS.get(i.get("caminho"), i.get("caminho")),
         i.get("segundos", 0), i.get("espera", 0), i.get("motivo", ""), s.get("id", "")]
        for s in sessoes for numero, i in enumerate(s.get("itens", []), 1)
    ]
    tabela(livro.create_sheet("Itens"), ["Data", "Computador", "Usuário", "Nº na sessão",
                                         "Situação", "Tipo", "Automação (s)", "Espera (s)",
                                         "Motivo do erro", "ID da sessão"], itens)

    # ---- Tempos manuais
    tabela(livro.create_sheet("Tempos manuais"), ["Data", "Hora", "Computador", "Usuário",
                                                  "Tipo", "Segundos", "Minutos"], [
        [m.get("data", "")[:10], m.get("data", "")[11:19], m.get("computador", ""),
         m.get("usuario", ""), ROTULO_CAMINHO.get(m.get("caminho"), m.get("caminho")),
         m.get("segundos", 0), _minutos(m.get("segundos", 0))] for m in manuais
    ])

    # ---- Agrupamentos
    for nome, chave in (("Por computador", "computador"), ("Por dia", "dia")):
        linhas = [[g[chave], g["sessoes"], g["feitos"], g["ja"], g["erros"],
                   round(g["segundos_automacao"], 1), round(g["segundos_manual"], 1),
                   round(g["segundos_economizados"], 1), _minutos(g["segundos_economizados"])]
                  for g in (ind.por_computador if chave == "computador" else ind.por_dia)]
        tabela(livro.create_sheet(nome), [chave.capitalize(), "Sessões", "Cadastrados",
                                          "Já tinham", "Erros", "Automação (s)",
                                          "Manual equivalente (s)", "Economia (s)",
                                          "Economia (min)"], linhas)

    tabela(livro.create_sheet("Motivos de erro"), ["Motivo", "Quantidade"],
           [list(m) for m in ind.motivos])

    metodologia = livro.create_sheet("Metodologia")
    metodologia["A1"] = "Como os números foram calculados"
    metodologia["A1"].font = titulo_fonte
    metodologia.column_dimensions["A"].width = 110
    for linha, texto in enumerate(METODOLOGIA, 3):
        metodologia.cell(row=linha, column=1, value=f"• {texto}").alignment = Alignment(wrap_text=True)

    livro.save(destino)


# ================================================================ HTML

def _barra(rotulo: str, segundos: float, maximo: float, cor: str) -> str:
    largura = 0 if not maximo else max(segundos / maximo * 100, 0.8)
    return (f'<div class="barra"><span class="rot">{html.escape(rotulo)}</span>'
            f'<span class="trilho"><span class="cheio" style="width:{largura:.1f}%;'
            f'background:{cor}"></span></span>'
            f'<span class="val">{formatar_duracao(segundos)}</span></div>')


def _gerar_html(destino: Path, ind: est.Indicadores, escopo: str) -> None:
    maximo = max(ind.media_manual_geral, ind.media_automacao)
    maximo_total = max(ind.segundos_manual_equivalente, ind.segundos_automacao)
    linhas = "".join(
        f'<tr><td{" class=sub" if str(a).startswith(" ") else ""}>{html.escape(str(a).strip())}</td>'
        f"<td>{html.escape(str(b))}</td></tr>" for a, b in _indicadores_tabela(ind))
    maquinas = "".join(
        f"<tr><td>{html.escape(g['computador'])}</td><td>{g['sessoes']}</td><td>{g['feitos']}</td>"
        f"<td>{g['erros']}</td><td>{formatar_duracao(g['segundos_economizados'])}</td></tr>"
        for g in ind.por_computador)
    motivos = "".join(f"<li>{html.escape(m)} — {n}</li>" for m, n in ind.motivos[:10]) or "<li>Nenhum</li>"
    metodologia = "".join(f"<li>{html.escape(t)}</li>" for t in METODOLOGIA)
    abrangencia = "todos os computadores" if escopo == est.TODOS else "este computador"

    pagina = f"""<!doctype html>
<html lang="pt-BR"><head><meta charset="utf-8">
<title>Relatório — Cadastro de Credores no SIGEF</title>
<style>
  :root {{ --tinta:#14213D; --suave:#3D4B66; --azul:#0B3D91; --verde:#1F7A3A; --amarelo:#FFD84D;
          --fundo:#F3F6FA; --borda:#C9D2E0; }}
  * {{ box-sizing:border-box; }}
  body {{ margin:0; background:var(--fundo); color:var(--tinta);
         font:16px/1.5 "Atkinson Hyperlegible","Segoe UI",Arial,sans-serif; }}
  header {{ background:var(--azul); color:#fff; padding:28px 48px; border-bottom:6px solid var(--amarelo); }}
  header h1 {{ margin:0; font-size:28px; }} header p {{ margin:6px 0 0; opacity:.9; }}
  main {{ max-width:1000px; margin:0 auto; padding:32px 48px 64px; }}
  h2 {{ font-size:21px; margin:36px 0 12px; }}
  .destaque {{ display:grid; grid-template-columns:repeat(4,1fr); gap:14px; }}
  .bloco {{ background:#fff; border:1px solid var(--borda); padding:16px 18px; }}
  .bloco b {{ display:block; font-size:30px; line-height:1.15; color:var(--azul); }}
  .bloco.verde b {{ color:var(--verde); }}
  .barra {{ display:grid; grid-template-columns:170px 1fr 130px; gap:12px; align-items:center; margin:8px 0; }}
  .trilho {{ height:26px; background:#E7EBF1; display:block; }}
  .cheio {{ height:100%; display:block; }}
  .val {{ font-weight:bold; }}
  table {{ width:100%; border-collapse:collapse; background:#fff; }}
  td, th {{ text-align:left; padding:8px 12px; border-bottom:1px solid var(--borda); }}
  th {{ background:#E7EBF1; }}
  td.sub {{ padding-left:36px; color:var(--suave); }}
  .nota {{ color:var(--suave); font-size:14px; }}
  @media (max-width:760px) {{ .destaque {{ grid-template-columns:1fr 1fr; }} main, header {{ padding-left:16px; padding-right:16px; }} }}
  @media print {{ body {{ background:#fff; }} header {{ -webkit-print-color-adjust:exact; print-color-adjust:exact; }}
                 .cheio {{ -webkit-print-color-adjust:exact; print-color-adjust:exact; }} }}
</style></head><body>
<header><h1>Cadastro de Credores no SIGEF — resultados da automação</h1>
<p>Dados de {abrangencia} · {formatar_data(ind.primeira_data)} a {formatar_data(ind.ultima_data)} ·
gerado em {datetime.now():%d/%m/%Y %H:%M}</p></header>
<main>
<div class="destaque">
  <div class="bloco"><b>{ind.feitos}</b>cadastros concluídos</div>
  <div class="bloco verde"><b>{formatar_duracao(ind.segundos_economizados)}</b>de trabalho economizado</div>
  <div class="bloco"><b>{formatar_vezes(ind.vezes_mais_rapido)}</b>mais rápido que o manual</div>
  <div class="bloco"><b>{ind.taxa_sucesso:.0%}</b>de sucesso</div>
</div>

<h2>Tempo por cadastro</h2>
{_barra("Manual", ind.media_manual_geral, maximo, "#8C9AB0")}
{_barra("Com a automação", ind.media_automacao, maximo, "var(--verde)")}

<h2>Tempo total do período</h2>
{_barra("Manual (equivalente)", ind.segundos_manual_equivalente, maximo_total, "#8C9AB0")}
{_barra("Com a automação", ind.segundos_automacao, maximo_total, "var(--verde)")}
<p class="nota">Tempo manual: {ind.fonte_manual(NOVO)} para credor novo; {ind.fonte_manual(EXISTENTE)} para só a conta.</p>

<h2>Indicadores</h2>
<table>{linhas}</table>

<h2>Por computador</h2>
<table><tr><th>Computador</th><th>Sessões</th><th>Cadastrados</th><th>Erros</th><th>Economia</th></tr>
{maquinas or '<tr><td colspan="5">Sem dados</td></tr>'}</table>

<h2>Motivos de erro mais comuns</h2>
<ul>{motivos}</ul>

<h2>Como os números foram calculados</h2>
<ul class="nota">{metodologia}</ul>
</main></body></html>"""
    destino.write_text(pagina, encoding="utf-8")


# ================================================================ geral

def gerar(escopo: str = est.LOCAL, pasta: Path = None) -> Path:
    """Gera o relatório e devolve a pasta criada."""
    sessoes, manuais = est.carregar(escopo)
    ind = est.calcular(sessoes, manuais, est.ler_ajustes()["estimativa"])

    nome = f"Relatorio_{datetime.now():%Y-%m-%d_%H%M%S}"
    destino = Path(pasta or pasta_padrao()) / nome

    try:
        destino.mkdir(parents=True, exist_ok=True)
    except OSError:
        # Documentos indisponível (perfil de rede, permissões): usa a pasta do programa
        destino = est.PASTA_ESTATISTICAS.parent / "relatorios" / nome
        destino.mkdir(parents=True, exist_ok=True)

    _gerar_excel(destino / "dados_pesquisa.xlsx", ind, sessoes, manuais)
    _gerar_html(destino / "relatorio.html", ind, escopo)

    with open(destino / "dados_brutos.jsonl", "w", encoding="utf-8") as arquivo:
        for registro in sessoes + manuais:
            arquivo.write(json.dumps(registro, ensure_ascii=False) + "\n")

    return destino
